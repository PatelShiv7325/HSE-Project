from datetime import datetime, date
from flask import Blueprint, render_template, request, redirect, url_for, flash, make_response
from flask_login import login_required, current_user
from app import db
from app.models import InspectionReport, Company
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from io import BytesIO
import re

inspections_bp = Blueprint("inspections", __name__, url_prefix="/inspections")

# Every field that lives inside InspectionReport.data for a Form 9
# (Hoists/Lifts examination certificate). report_no and report_date are
# stored as real columns (for search/sort/uniqueness); everything else
# lives in this JSON blob. Add new keys here to add a field to the form --
# no DB migration needed since `data` is a single JSON column.
FORM9_FIELDS = [
    "reg_no", "license_no", "nic_code", "occupier_name", "address",
    "hoist_description", "hoist_make", "tag_no", "capacity", "no_of_floors",
    "location", "construction_date",
    "good_mechanical_order", "enclosure_findings", "landing_gates_findings",
    "interlock_findings", "other_gate_fastenings", "cage_platform_findings",
    "overrunning_devices", "suspension_ropes", "safety_gear", "brakes",
    "worm_gearing", "electrical_equipment", "other_parts",
    "inaccessible_parts", "max_safe_load", "repairs_required", "remarks",
    "certification_date", "next_exam_date", "reminder_date",
    "competent_person_name", "competent_person_title",
    "competent_person_authority", "competent_person_no",
    "competent_person_state", "certification_text",
]

# Fields on the Form 9 that default to "Satisfactory" when left blank --
# these are the pass/fail findings for each part inspected.
FORM9_FINDING_FIELDS = [
    "good_mechanical_order", "enclosure_findings", "landing_gates_findings",
    "interlock_findings", "other_gate_fastenings", "cage_platform_findings",
    "overrunning_devices", "suspension_ropes", "safety_gear", "brakes",
    "worm_gearing", "electrical_equipment", "other_parts",
]


def _parse_date(value):
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d").date()


@inspections_bp.route("/form9")
@login_required
def form9_list():
    search = request.args.get("q", "").strip()
    per_page = request.args.get("per_page", 10, type=int)
    page = request.args.get("page", 1, type=int)

    query = InspectionReport.query.filter_by(form_type="form9")
    if search:
        query = query.filter(
            db.or_(
                InspectionReport.report_no.ilike(f"%{search}%"),
                InspectionReport.occupier_name.ilike(f"%{search}%"),
            )
        )
    query = query.order_by(InspectionReport.report_date.desc().nullslast(), InspectionReport.id.desc())

    total = query.count()
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    reports = query.offset((page - 1) * per_page).limit(per_page).all()

    start = 0 if total == 0 else (page - 1) * per_page + 1
    end = min(page * per_page, total)

    return render_template(
        "inspections/form9_list.html",
        reports=reports, search=search, per_page=per_page,
        page=page, total_pages=total_pages, total=total, start=start, end=end,
    )


@inspections_bp.route("/form9/new", methods=["GET", "POST"])
@login_required
def form9_new():
    return _form9_save(report=None)


@inspections_bp.route("/form9/<int:report_id>/edit", methods=["GET", "POST"])
@login_required
def form9_edit(report_id):
    report = InspectionReport.query.filter_by(id=report_id, form_type="form9").first_or_404()
    return _form9_save(report=report)


def _form9_save(report):
    if request.method == "POST":
        report_no = request.form.get("report_no", "").strip()
        if not report_no:
            flash("Report No. is required.", "danger")
            return redirect(request.url)

        existing = InspectionReport.query.filter(
            InspectionReport.report_no == report_no,
            InspectionReport.form_type == "form9",
        )
        if report:
            existing = existing.filter(InspectionReport.id != report.id)
        if existing.first():
            flash("A Form 9 report with this Report No. already exists.", "danger")
            return redirect(request.url)

        data = {}
        for field in FORM9_FIELDS:
            value = request.form.get(field, "").strip()
            if not value and field in FORM9_FINDING_FIELDS:
                value = "Satisfactory"
            data[field] = value

        if report is None:
            report = InspectionReport(form_type="form9", created_by_id=current_user.id)
            db.session.add(report)

        company_id = request.form.get("company_id", type=int)
        report.company_id = company_id or None

        report.report_no = report_no
        report.report_date = _parse_date(request.form.get("date")) or date.today()
        report.occupier_name = data.get("occupier_name")
        report.data = data
        db.session.commit()

        flash("Form 9 report saved.", "success")
        return redirect(url_for("inspections.form9_list"))

    return render_template(
        "inspections/form9_form.html",
        report=report,
        data=(report.data if report else {}),
        today=date.today().isoformat(),
        companies=Company.query.order_by(Company.name.asc()).all(),
    )


@inspections_bp.route("/form9/<int:report_id>/delete", methods=["POST"])
@login_required
def form9_delete(report_id):
    report = InspectionReport.query.filter_by(id=report_id, form_type="form9").first_or_404()
    db.session.delete(report)
    db.session.commit()
    flash("Form 9 report deleted.", "success")
    return redirect(url_for("inspections.form9_list"))


@inspections_bp.route("/form9/<int:report_id>/pdf")
@login_required
def form9_pdf(report_id):
    report = InspectionReport.query.filter_by(id=report_id, form_type="form9").first_or_404()

    # WeasyPrint is an optional dependency -- imported here (not at module
    # level) so the rest of the app still runs even before `pip install`
    # has been re-run to pick it up from requirements.txt.
    try:
        from weasyprint import HTML
    except ImportError:
        flash("PDF export needs the 'weasyprint' package -- run: pip install -r requirements.txt", "danger")
        return redirect(url_for("inspections.form9_list"))

    html = render_template("inspections/form9_pdf.html", report=report, data=report.data)
    pdf_bytes = HTML(string=html, base_url=request.url_root).write_pdf()

    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f'inline; filename="Form9-{report.report_no}.pdf"'
    return response


# ---------------------------------------------------------------------------
# Generic engine for Form 10, Form 11, PSV and Centrifuge.
#
# Form 9 above is hand-built (one route per action). These four forms share
# the exact same InspectionReport/JSON pattern, so instead of copy-pasting
# Form 9's routes four more times, every field for every one of these forms
# is declared once in FORMS_CONFIG below, and a single set of routes
# (parameterized by <form_type>) drives list/new/edit/delete/pdf/renew for
# whichever of the four the URL asks for. Add a fifth form the same way:
# just add one more entry to FORMS_CONFIG and to the `any(...)` converter
# in each @inspections_bp.route() below -- no new view functions needed.
# ---------------------------------------------------------------------------

FORMS_CONFIG = {
    "form10": {
        "label": "Form 10 - Lifting Machines / Cranes",
        "rule": "(Prescribed under Rule 60)",
        "sections": [
            ("Registration", [
                ("reg_no", "Registration No.", "text", None),
                ("license_no", "License No.", "text", None),
                ("nic_code", "NIC Code", "text", None),
            ]),
            ("Equipment", [
                ("equipment_description", "Equipment Description", "text", None),
                ("serial_no", "Serial No.", "text", None),
                ("make", "Make", "text", None),
                ("capacity", "Capacity", "text", None),
                ("location", "Location", "text", None),
            ]),
            ("Dates", [
                ("first_use_date", "First Use Date", "date", None),
                ("last_exam_date", "Last Exam Date", "date", None),
                ("examination_date", "Examination Date", "date", None),
                ("examined_by", "Examined By", "text", None),
            ]),
            ("Findings", [
                ("examination_details_text", "Examination Details", "textarea", None),
                ("certificate_details", "Certificate Details", "textarea", None),
                ("heat_treatment_details", "Heat Treatment Details", "textarea", None),
                ("defects_found", "Defects Found", "textarea", None),
            ]),
            ("Certification", [
                ("certification_text", "Certification Text", "textarea",
                 "I certify that on {date} the lifting machine, chain, rope or lifting tackle described above was thoroughly examined and that the above is a true report of the result of the examination."),
                ("certification_date", "Certification Date", "date", None),
                ("next_exam_date", "Next Exam Date", "date", None),
                ("reminder_date", "Reminder Date", "date", None),
                ("competent_person_name", "Competent Person Name", "text", None),
                ("competent_person_title", "Title", "text", "Competent Person Declared by Director"),
                ("competent_person_authority", "Authority", "text", "Industrial Safety & Health Gujarat State"),
                ("competent_person_no", "Competent Person No.", "text", None),
                ("competent_person_state", "State", "text", "Gujarat"),
            ]),
        ],
        "finding_fields": [],
    },
    "form11": {
        "label": "Form 11 - Pressure Vessel",
        "rule": "(Prescribed under Rule 61 of GFR 1963)",
        "sections": [
            ("Registration", [
                ("reg_no", "Registration No.", "text", None),
                ("license_no", "License No.", "text", None),
                ("nic_code", "NIC Code", "text", None),
            ]),
            ("Vessel Details", [
                ("vessel_name", "Vessel Name", "text", None),
                ("vessel_description", "Vessel Description", "text", None),
                ("tag_no", "Tag No.", "text", None),
                ("capacity", "Capacity", "text", None),
                ("location", "Location", "text", None),
                ("manufacturer", "Manufacturer", "text", None),
                ("nature_of_process", "Nature of Process", "text", None),
                ("temperature", "Temperature", "text", None),
                ("pressure", "Pressure", "text", None),
                ("date_of_construction", "Date of Construction", "date", None),
                ("safe_working_pressure", "Safe Working Pressure", "text", None),
            ]),
            ("Thickness", [
                ("thickness_shell", "Shell Thickness", "text", None),
                ("thickness_jacket", "Jacket Thickness", "text", None),
                ("thickness_limpet", "Limpet Thickness", "text", None),
                ("thickness_pipeline", "Pipeline Thickness", "text", None),
            ]),
            ("Examination History", [
                ("first_use_date", "First Use Date", "date", None),
                ("last_exam_date", "Last Exam Date", "date", None),
                ("last_external_exam", "Last External Exam", "text", None),
                ("last_internal_exam", "Last Internal Exam", "text", "Not Applicable"),
                ("last_hydraulic_exam", "Last Hydraulic Exam", "text", None),
                ("last_hydraulic_exam_date", "Last Hydraulic Exam Date", "date", None),
                ("last_hydraulic_exam_comment", "Last Hydraulic Exam Comment", "text", None),
                ("last_ultrasonic_test", "Last Ultrasonic Test", "text", None),
                ("last_ultrasonic_test_date", "Last Ultrasonic Test Date", "date", None),
                ("last_ultrasonic_test_comment", "Last Ultrasonic Test Comment", "text", None),
                ("last_hydro_test_date", "Last Hydro Test Date", "date", None),
            ]),
            ("Findings", [
                ("lagging_removed", "Lagging Removed", "text", None),
                ("external_findings", "External Findings", "textarea", None),
                ("internal_findings", "Internal Findings", "textarea", None),
                ("ultrasonic_findings", "Ultrasonic Findings", "textarea", None),
                ("ultrasonic_shell", "Ultrasonic - Shell", "text", None),
                ("ultrasonic_jacket", "Ultrasonic - Jacket", "text", None),
                ("ultrasonic_limpet", "Ultrasonic - Limpet", "text", None),
                ("ultrasonic_pipeline", "Ultrasonic - Pipeline", "text", None),
                ("vessel_condition", "Vessel Condition", "text", None),
                ("piping_condition", "Piping Condition", "text", None),
                ("pressure_gauges_condition", "Pressure Gauges Condition", "text", None),
                ("safety_valve_condition", "Safety Valve Condition", "text", None),
                ("stop_valve_condition", "Stop Valve Condition", "text", None),
                ("reducing_valve_condition", "Reducing Valve Condition", "text", None),
                ("additional_safety_valve_condition", "Additional Safety Valve Condition", "text", None),
                ("other_devices_condition", "Other Devices Condition", "text", None),
            ]),
            ("Repairs & Safety", [
                ("repairs_required", "Repairs Required", "textarea", None),
                ("repairs_period", "Repairs Period", "text", None),
                ("safety_conditions", "Safety Conditions", "textarea",
                 "Client is advised for regular Testing of Safety Valve & Keeping Records Every Year."),
                ("reduced_working_pressure", "Reduced Working Pressure", "text", "Not Applicable"),
                ("calculated_swp", "Calculated SWP", "text", None),
                ("recommended_swp", "Recommended SWP", "text", None),
                ("other_observations", "Other Observations", "textarea",
                 "Pressure Gauge and Safety Valve should be checked periodically to ensure correct functioning and safe operation. Proper inspection and maintenance records should be maintained for compliance and safety purposes."),
            ]),
            ("Certification", [
                ("certification_text", "Certification Text", "textarea",
                 "I certify that on {date} the pressure vessel or plant described above was thorough cleaned and (so far, its construction permits) made accessible for thorough examination and for such tests as were necessary for thorough examination and that on the said date I thoroughly examination this pressure vessel or plants, including its fitting and that above is a true report of examination."),
                ("certification_date", "Certification Date", "date", None),
                ("next_ndt_date", "Next NDT Date", "date", None),
                ("next_hydro_date", "Next Hydro Date", "date", None),
                ("next_exam_date", "Next Exam Date", "date", None),
                ("reminder_date", "Reminder Date", "date", None),
                ("competent_person_name", "Competent Person Name", "text", None),
                ("competent_person_title", "Title", "text", "Competent Person Declared by Director"),
                ("competent_person_authority", "Authority", "text", "Industrial Safety & Health Gujarat State"),
                ("competent_person_no", "Competent Person No.", "text", None),
                ("competent_person_state", "State", "text", "Gujarat"),
            ]),
        ],
        "finding_fields": [
            "vessel_condition", "piping_condition", "pressure_gauges_condition",
            "safety_valve_condition", "stop_valve_condition", "reducing_valve_condition",
            "additional_safety_valve_condition", "other_devices_condition",
        ],
    },
    "psv": {
        "label": "PSV - Pressure Safety Valve Certificate",
        "rule": "Pressure Safety Valve Test Certificate",
        "sections": [
            ("Valve Details", [
                ("fitted_location", "Fitted Location", "text", None),
                ("year_of_mfg", "Year of Manufacture", "text", None),
                ("make", "Make", "text", None),
                ("humidity", "Humidity", "text", None),
                ("set_pressure", "Set Pressure", "text", None),
                ("temperature", "Temperature", "text", None),
            ]),
            ("Testing Standard & Traceability", [
                ("cal_std_used", "Calibration Standard Used", "text", None),
                ("cal_accuracy", "Calibration Accuracy", "text", None),
                ("cal_make", "Calibration Equipment Make", "text", None),
                ("cal_by", "Calibrated By", "text", None),
                ("cal_equip_used", "Calibration Equipment Used", "text", None),
                ("cal_cert_no", "Calibration Certificate No.", "text", None),
                ("cal_model", "Calibration Equipment Model", "text", None),
                ("cal_date", "Calibration Date", "date", None),
                ("cal_sr_no", "Calibration Equipment Sr. No.", "text", None),
                ("cal_due", "Calibration Due Date", "date", None),
                ("cal_range", "Calibration Range", "text", None),
                ("cal_least_count", "Least Count", "text", None),
                ("cal_traceability", "Traceability", "text", None),
                ("operated_at", "Operated At", "text", None),
            ]),
            ("Certification", [
                ("next_exam_date", "Next Exam Date", "date", None),
                ("reminder_date", "Reminder Date", "date", None),
                ("competent_person_name", "Competent Person Name", "text", None),
                ("competent_person_no", "Competent Person No.", "text", None),
                ("competent_person_state", "State", "text", None),
                ("competent_person_title", "Title", "text", "CHARTER ENGINEER"),
            ]),
        ],
        "finding_fields": [],
    },
    "centrifuge": {
        "label": "Centrifuge Machine Test Report",
        "rule": "Centrifuge Machine Test Report",
        "sections": [
            ("Registration", [
                ("license_no", "License No.", "text", None),
            ]),
            ("Machine Details", [
                ("machine_name_description", "Machine Name/Description", "text", None),
                ("machine_name", "Machine Name", "text", None),
                ("tag_no", "Tag No.", "text", None),
                ("capacity", "Capacity", "text", None),
                ("location", "Location", "text", None),
                ("manufacturer_name_address", "Manufacturer Name & Address", "textarea", None),
                ("date_of_construction", "Date of Construction", "date", None),
                ("machine_size", "Machine Size", "text", None),
            ]),
            ("Safety Checks", [
                ("condition_of_machine", "Condition of Machine", "text", "Satisfactory"),
                ("interlock_top_cover", "Interlock - Top Cover", "text", "Satisfactory"),
                ("interlock_mechanical_breaker", "Interlock - Mechanical Breaker", "text", "Satisfactory"),
                ("earthing_provided", "Earthing Provided", "text", "Satisfactory"),
                ("basket_speed", "Basket Speed", "text", None),
                ("operating_speed_stamped", "Operating Speed (Stamped)", "text", None),
            ]),
            ("Examination", [
                ("last_exam_date", "Last Exam Date", "date", None),
                ("remarks", "Remarks", "textarea", None),
                ("date_of_examination", "Date of Examination", "date", None),
                ("defects_and_remedies", "Defects & Remedies", "textarea", None),
            ]),
            ("Certification", [
                ("certification_date", "Certification Date", "date", None),
                ("certification_text", "Certification Text", "textarea",
                 "I certify that on {date} I have thoroughly examined the centrifuge machine described above and the above is a correct report of the result of such examination."),
                ("next_exam_date", "Next Exam Date", "date", None),
                ("next_ndt_date", "Next NDT Date", "date", None),
                ("next_hydro_date", "Next Hydro Date", "date", None),
                ("reminder_date", "Reminder Date", "date", None),
                ("competent_person_name", "Competent Person Name", "text", None),
                ("competent_person_no", "Competent Person No.", "text", None),
                ("competent_person_state", "State", "text", "Gujarat"),
                ("competent_person_title", "Title", "text",
                 "Competent Person Declared by Director Industrial Safety & Health"),
            ]),
        ],
        "finding_fields": [
            "condition_of_machine", "interlock_top_cover",
            "interlock_mechanical_breaker", "earthing_provided",
        ],
    },
}


def _fields_for(form_type):
    """Flat list of (key, label, input_type, default) across every section."""
    fields = []
    for _section_name, section_fields in FORMS_CONFIG[form_type]["sections"]:
        fields.extend(section_fields)
    return fields


def _suggest_report_no(form_type):
    """FORM10/2026/0001-style suggestion, just a prefill -- the user can
    still change it before saving, and _generic_save re-checks uniqueness."""
    from datetime import date as _date
    year = _date.today().year
    count = InspectionReport.query.filter(
        InspectionReport.form_type == form_type,
        db.extract("year", InspectionReport.report_date) == year,
    ).count()
    prefix = form_type.upper()
    return f"{prefix}/{year}/{count + 1:04d}"


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/")
@login_required
def generic_list(form_type):
    search = request.args.get("q", "").strip()
    per_page = request.args.get("per_page", 10, type=int)
    page = request.args.get("page", 1, type=int)

    query = InspectionReport.query.filter_by(form_type=form_type)
    if search:
        query = query.filter(
            db.or_(
                InspectionReport.report_no.ilike(f"%{search}%"),
                InspectionReport.occupier_name.ilike(f"%{search}%"),
            )
        )
    query = query.order_by(InspectionReport.report_date.desc().nullslast(), InspectionReport.id.desc())

    total = query.count()
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    reports = query.offset((page - 1) * per_page).limit(per_page).all()

    start = 0 if total == 0 else (page - 1) * per_page + 1
    end = min(page * per_page, total)

    return render_template(
        "inspections/generic_list.html",
        form_type=form_type, config=FORMS_CONFIG[form_type],
        reports=reports, search=search, per_page=per_page,
        page=page, total_pages=total_pages, total=total, start=start, end=end,
    )


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/new", methods=["GET", "POST"])
@login_required
def generic_new(form_type):
    return _generic_save(form_type, report=None)


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/<int:report_id>/edit", methods=["GET", "POST"])
@login_required
def generic_edit(form_type, report_id):
    report = InspectionReport.query.filter_by(id=report_id, form_type=form_type).first_or_404()
    return _generic_save(form_type, report=report)


def _generic_save(form_type, report):
    config = FORMS_CONFIG[form_type]
    fields = _fields_for(form_type)
    finding_fields = set(config["finding_fields"])

    if request.method == "POST":
        report_no = request.form.get("report_no", "").strip()
        if not report_no:
            flash("Report No. is required.", "danger")
            return redirect(request.url)

        existing = InspectionReport.query.filter(
            InspectionReport.report_no == report_no,
            InspectionReport.form_type == form_type,
        )
        if report:
            existing = existing.filter(InspectionReport.id != report.id)
        if existing.first():
            flash(f"A {config['label']} report with this Report No. already exists.", "danger")
            return redirect(request.url)

        data = {}
        for key, _label, _input_type, default in fields:
            value = request.form.get(key, "").strip()
            if not value and key in finding_fields:
                value = "Satisfactory"
            if not value and default:
                value = default
            data[key] = value

        # occupier_name / address are common to every form, entered once at
        # the top of the generic form template (same as Form 9), not part
        # of FORMS_CONFIG's per-form field list.
        data["occupier_name"] = request.form.get("occupier_name", "").strip()
        data["address"] = request.form.get("address", "").strip()

        if report is None:
            report = InspectionReport(form_type=form_type, created_by_id=current_user.id)
            db.session.add(report)

        company_id = request.form.get("company_id", type=int)
        report.company_id = company_id or None
        report.report_no = report_no
        report.report_date = _parse_date(request.form.get("date")) or date.today()
        report.occupier_name = data.get("occupier_name")
        report.data = data
        db.session.commit()

        flash(f"{config['label']} report saved.", "success")
        return redirect(url_for("inspections.generic_list", form_type=form_type))

    return render_template(
        "inspections/generic_form.html",
        form_type=form_type, config=config,
        report=report, data=(report.data if report else {}),
        today=date.today().isoformat(),
        suggested_report_no=(report.report_no if report else _suggest_report_no(form_type)),
        companies=Company.query.order_by(Company.name.asc()).all(),
    )


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/<int:report_id>/delete", methods=["POST"])
@login_required
def generic_delete(form_type, report_id):
    report = InspectionReport.query.filter_by(id=report_id, form_type=form_type).first_or_404()
    db.session.delete(report)
    db.session.commit()
    flash(f"{FORMS_CONFIG[form_type]['label']} report deleted.", "success")
    return redirect(url_for("inspections.generic_list", form_type=form_type))


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/<int:report_id>/renew", methods=["POST"])
@login_required
def generic_renew(form_type, report_id):
    """Copies an existing report's data into a brand-new report (new Report
    No., new id), links it back via renewed_from_id, then sends the user
    straight to editing the copy so they can update dates/findings before
    saving it for real."""
    original = InspectionReport.query.filter_by(id=report_id, form_type=form_type).first_or_404()

    new_report = InspectionReport(
        form_type=form_type,
        created_by_id=current_user.id,
        report_no=_suggest_report_no(form_type),
        report_date=date.today(),
        occupier_name=original.occupier_name,
        company_id=original.company_id,
        renewed_from_id=original.id,
        data=dict(original.data or {}),
    )
    db.session.add(new_report)
    db.session.commit()

    flash(f"Created renewal {new_report.report_no} from {original.report_no}. Review and save.", "success")
    return redirect(url_for("inspections.generic_edit", form_type=form_type, report_id=new_report.id))


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/<int:report_id>/pdf")
@login_required
def generic_pdf(form_type, report_id):
    report = InspectionReport.query.filter_by(id=report_id, form_type=form_type).first_or_404()
    config = FORMS_CONFIG[form_type]

    try:
        from weasyprint import HTML
    except ImportError:
        flash("PDF export needs the 'weasyprint' package -- run: pip install -r requirements.txt", "danger")
        return redirect(url_for("inspections.generic_list", form_type=form_type))

    html = render_template(
        "inspections/generic_pdf.html",
        report=report, data=report.data, config=config, form_type=form_type,
    )
    pdf_bytes = HTML(string=html, base_url=request.url_root).write_pdf()

    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f'inline; filename="{form_type.upper()}-{report.report_no}.pdf"'
    return response


# ---------------------------------------------------------------------------
# Phase 3 -- Compliance Dashboard: expiry tracker across every form type
# (Form 9/10/11, PSV, Centrifuge), a 6-month volume chart, and a recent
# activity feed. Pure read-only aggregation over InspectionReport -- no new
# DB columns, no schema changes needed.
# ---------------------------------------------------------------------------

# Which InspectionReport.data key(s) hold the "next due" date for each form
# type, in priority order (first one present wins) -- mirrors the fields
# actually offered in FORMS_CONFIG / FORM9_FIELDS above.
DUE_DATE_KEYS = {
    "form9": ["next_exam_date"],
    "form10": ["next_exam_date"],
    "form11": ["next_exam_date", "next_ndt_date", "next_hydro_date"],
    "psv": ["next_exam_date"],
    "centrifuge": ["next_exam_date", "next_ndt_date", "next_hydro_date"],
}

FORM_LABELS = {
    "form9": "Form 9", "form10": "Form 10", "form11": "Form 11",
    "psv": "PSV", "centrifuge": "Centrifuge",
}


def _report_due_date(report):
    """First present due-date field for this report's form type, parsed."""
    for key in DUE_DATE_KEYS.get(report.form_type, []):
        raw = (report.data or {}).get(key)
        if raw:
            try:
                return _parse_date(raw)
            except (ValueError, TypeError):
                continue
    return None


@inspections_bp.route("/dashboard")
@login_required
def dashboard():
    reports = InspectionReport.query.all()
    today = date.today()

    # A report counts as "Renewed" once a newer report points back at it
    # via renewed_from_id -- no report_no regex needed, the DB already knows.
    renewed_source_ids = {r.renewed_from_id for r in reports if r.renewed_from_id}

    rows = []
    for r in reports:
        due_date = _report_due_date(r)
        days_left = (due_date - today).days if due_date else None

        if r.id in renewed_source_ids:
            status = "renewed"
        elif days_left is None:
            status = "valid"
        elif days_left < 0:
            status = "expired"
        elif days_left < 15:
            status = "critical"
        elif days_left < 30:
            status = "reminder"
        else:
            status = "valid"

        rows.append({
            "report": r,
            "form_type": r.form_type,
            "form_label": FORM_LABELS.get(r.form_type, r.form_type),
            "due_date": due_date,
            "days_left": days_left,
            "status": status,
            "company_name": (r.company.name if r.company_id and r.company else None) or r.occupier_name,
        })

    stats = {
        "total_companies": Company.query.count(),
        "total_reports": len(rows),
        "critical": sum(1 for x in rows if x["status"] == "critical"),
        "reminder": sum(1 for x in rows if x["status"] == "reminder"),
        "valid": sum(1 for x in rows if x["status"] in ("valid", "renewed")),
        "expired": sum(1 for x in rows if x["status"] == "expired"),
    }

    # Expiry tracker table: anything due within 30 days (or already expired),
    # soonest first. Renewed/valid rows are reachable via the filter pills
    # in the template (client-side, no reload needed).
    expiring = sorted(
        [x for x in rows if x["days_left"] is not None and x["days_left"] <= 30],
        key=lambda x: x["days_left"],
    )

    # Recent activity: last 20 touched reports, newest first.
    recent = sorted(
        rows,
        key=lambda x: x["report"].updated_at or x["report"].created_at or datetime.min,
        reverse=True,
    )[:20]
    for x in recent:
        created = x["report"].created_at
        updated = x["report"].updated_at
        x["edited"] = bool(created and updated and (updated - created).total_seconds() > 2)

    # 6-month report volume, oldest -> newest, by created_at.
    months = []
    y, m = today.year, today.month
    for _ in range(6):
        months.insert(0, {"year": y, "month": m, "label": date(y, m, 1).strftime("%b"), "count": 0})
        m -= 1
        if m == 0:
            m, y = 12, y - 1
    for x in rows:
        created = x["report"].created_at
        if not created:
            continue
        for bucket in months:
            if created.year == bucket["year"] and created.month == bucket["month"]:
                bucket["count"] += 1
                break
    max_count = max([b["count"] for b in months] + [1])
    for b in months:
        b["pct"] = round((b["count"] / max_count) * 100)

    return render_template(
        "inspections/dashboard.html",
        stats=stats,
        expiring=expiring,
        recent=recent,
        chart_months=months,
    )


# ---------------------------------------------------------------------------
# Phase 4 -- Bulk PDF export (merge every matching report into one PDF, one
# page per report) and Excel registry export (branded .xlsx), for Form 9
# and the generic four (Form 10/11, PSV, Centrifuge). Both respect the same
# "q" search box already used on each list page, so "export what I'm
# looking at" works exactly like it does on screen.
# ---------------------------------------------------------------------------

# Curated identification columns per form type for the Excel registry --
# deliberately a subset of FORMS_CONFIG (which has 20-40+ fields per form)
# so the exported sheet stays readable rather than dumping every field.
EXCEL_COLUMNS = {
    "form9": [
        ("Location", "location"),
        ("Hoist Description", "hoist_description"),
        ("Make", "hoist_make"),
        ("Tag No.", "tag_no"),
        ("Capacity", "capacity"),
    ],
    "form10": [
        ("Location", "location"),
        ("Equipment Description", "equipment_description"),
        ("Make", "make"),
        ("Capacity", "capacity"),
    ],
    "form11": [
        ("Vessel Name", "vessel_name"),
        ("Tag No.", "tag_no"),
        ("Capacity", "capacity"),
        ("Location", "location"),
        ("Safe Working Pressure", "safe_working_pressure"),
        ("Last Hydro Test Date", "last_hydro_test_date"),
    ],
    "psv": [
        ("Fitted Location", "fitted_location"),
        ("Make", "make"),
        ("Set Pressure", "set_pressure"),
        ("Calibration Due", "cal_due"),
    ],
    "centrifuge": [
        ("Machine Name", "machine_name"),
        ("Tag No.", "tag_no"),
        ("Capacity", "capacity"),
        ("Location", "location"),
    ],
}


def _export_query(form_type, search):
    """Same filter used by the list pages, but unpaginated -- export acts on
    everything the search box currently matches, not just the visible page."""
    query = InspectionReport.query.filter_by(form_type=form_type)
    if search:
        query = query.filter(
            db.or_(
                InspectionReport.report_no.ilike(f"%{search}%"),
                InspectionReport.occupier_name.ilike(f"%{search}%"),
            )
        )
    return query.order_by(InspectionReport.report_no.asc()).all()


def _status_and_days(report, renewed_source_ids):
    """Same bucketing rules as the compliance dashboard (Phase 3) -- kept as
    a small standalone helper here so the registry's Status/Days columns
    always agree with what the dashboard shows."""
    due_date = _report_due_date(report)
    today = date.today()
    days_left = (due_date - today).days if due_date else None

    if report.id in renewed_source_ids:
        status = "Renewed"
    elif days_left is None:
        status = "Valid"
    elif days_left < 0:
        status = "Expired"
    elif days_left < 15:
        status = "Critical"
    elif days_left < 30:
        status = "Reminder"
    else:
        status = "Valid"
    return status, due_date, days_left


def _renewed_source_ids():
    return {
        r.renewed_from_id
        for r in InspectionReport.query.with_entities(InspectionReport.renewed_from_id).all()
        if r.renewed_from_id
    }


def _company_display_name(report):
    return (report.company.name if report.company_id and report.company else None) or report.occupier_name or "-"


def _build_excel_registry(form_type, reports, label):
    renewed_source_ids = _renewed_source_ids()
    extra_columns = EXCEL_COLUMNS.get(form_type, [])

    headers = ["Sr No.", "Report No.", "Report Date", "Company / Occupier"]
    headers += [col_label for col_label, _key in extra_columns]
    headers += ["Due Date", "Days Left", "Status"]

    wb = Workbook()
    ws = wb.active
    # Excel sheet names can't contain \ / ? * [ ] or exceed 31 chars --
    # form labels like "Form 10 - Lifting Machines / Cranes" have a slash,
    # so strip anything invalid rather than let openpyxl reject it.
    safe_title = re.sub(r'[\\/?*\[\]:]', '-', label)[:31]
    ws.title = safe_title

    col_count = len(headers)

    # Branding header rows (merged across the full table width), same idea
    # as the "COMPANY NAME / TAGLINE / REGISTRY TITLE" block used elsewhere.
    from flask import current_app
    company_name = current_app.config.get("COMPANY_NAME", "HSE Project")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    ws.cell(row=1, column=1, value=company_name.upper()).font = Font(bold=True, size=14)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    ws.cell(row=2, column=1, value=f"{label.upper()} REGISTRY").font = Font(bold=True, size=11)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_count)
    ws.cell(row=3, column=1, value=f"Generated {datetime.utcnow().strftime('%d-%m-%Y %H:%M')} UTC").font = Font(italic=True, size=9)
    ws.cell(row=3, column=1).alignment = Alignment(horizontal="center")

    header_row = 5
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    row_idx = header_row + 1
    for i, r in enumerate(reports, start=1):
        status, due_date, days_left = _status_and_days(r, renewed_source_ids)
        data = r.data or {}
        row = [
            i,
            r.report_no,
            r.report_date.strftime("%d-%m-%Y") if r.report_date else "",
            _company_display_name(r),
        ]
        row += [(data.get(key) or "NA") for _label, key in extra_columns]
        row += [
            due_date.strftime("%d-%m-%Y") if due_date else "NA",
            days_left if days_left is not None else "NA",
            status,
        ]
        for c, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=c, value=value)
        row_idx += 1

    # Auto-size columns based on header/content length.
    for c in range(1, col_count + 1):
        max_len = len(str(headers[c - 1]))
        for r_idx in range(header_row + 1, row_idx):
            v = ws.cell(row=r_idx, column=c).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(c)].width = max_len + 4

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@inspections_bp.route("/form9/export/pdf")
@login_required
def form9_export_pdf():
    search = request.args.get("q", "").strip()
    reports = _export_query("form9", search)
    if not reports:
        flash("No Form 9 reports match the current search to export.", "danger")
        return redirect(url_for("inspections.form9_list", q=search))

    try:
        from weasyprint import HTML
    except ImportError:
        flash("PDF export needs the 'weasyprint' package -- run: pip install -r requirements.txt", "danger")
        return redirect(url_for("inspections.form9_list", q=search))

    html = render_template("inspections/form9_bulk_pdf.html", reports=reports)
    pdf_bytes = HTML(string=html, base_url=request.url_root).write_pdf()

    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f'inline; filename="Form9-Registry-{date.today().isoformat()}.pdf"'
    return response


@inspections_bp.route("/form9/export/excel")
@login_required
def form9_export_excel():
    search = request.args.get("q", "").strip()
    reports = _export_query("form9", search)
    if not reports:
        flash("No Form 9 reports match the current search to export.", "danger")
        return redirect(url_for("inspections.form9_list", q=search))

    buffer = _build_excel_registry("form9", reports, "Form 9")
    response = make_response(buffer.read())
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response.headers["Content-Disposition"] = f'attachment; filename="Form9-Registry-{date.today().isoformat()}.xlsx"'
    return response


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/export/pdf")
@login_required
def generic_export_pdf(form_type):
    search = request.args.get("q", "").strip()
    reports = _export_query(form_type, search)
    config = FORMS_CONFIG[form_type]
    if not reports:
        flash(f"No {config['label']} reports match the current search to export.", "danger")
        return redirect(url_for("inspections.generic_list", form_type=form_type, q=search))

    try:
        from weasyprint import HTML
    except ImportError:
        flash("PDF export needs the 'weasyprint' package -- run: pip install -r requirements.txt", "danger")
        return redirect(url_for("inspections.generic_list", form_type=form_type, q=search))

    html = render_template("inspections/generic_bulk_pdf.html", reports=reports, config=config, form_type=form_type)
    pdf_bytes = HTML(string=html, base_url=request.url_root).write_pdf()

    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f'inline; filename="{form_type.upper()}-Registry-{date.today().isoformat()}.pdf"'
    return response


@inspections_bp.route("/<any(form10, form11, psv, centrifuge):form_type>/export/excel")
@login_required
def generic_export_excel(form_type):
    search = request.args.get("q", "").strip()
    reports = _export_query(form_type, search)
    config = FORMS_CONFIG[form_type]
    if not reports:
        flash(f"No {config['label']} reports match the current search to export.", "danger")
        return redirect(url_for("inspections.generic_list", form_type=form_type, q=search))

    buffer = _build_excel_registry(form_type, reports, config["label"])
    response = make_response(buffer.read())
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response.headers["Content-Disposition"] = f'attachment; filename="{form_type.upper()}-Registry-{date.today().isoformat()}.xlsx"'
    return response